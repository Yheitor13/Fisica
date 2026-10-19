from pathlib import Path
import json
import numpy as np
from scipy import odr
from scipy.optimize import least_squares
from openpyxl import load_workbook
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
base=Path(__file__).resolve().parent.parent
w=load_workbook(base/'Tabela_Medidas_Originais.xlsx',data_only=True)
rows=list(w['Medidas'].iter_rows(min_row=2,max_row=12,min_col=1,max_col=5,values_only=True))
a=np.array([r[:4] for r in rows[:10]],float);x=a[:,0]/100;ys=a[:,1:]/100
mean=ys.mean(axis=1);s=ys.std(axis=1,ddof=1);sem=s/np.sqrt(3);ux=np.full(10,.01);uy=np.sqrt(sem**2+.01**2)
X=np.log(x);Y=np.log(mean);sx=ux/x;sy=uy/mean;Q=x*x;uq=2*x*ux;g=9.79

def fit(xx,yy,sxx,syy,kind,beta):
 f=(lambda b,z:b[0]*z+b[1]) if kind=='line' else (lambda b,z:b[0]*z)
 o=odr.ODR(odr.RealData(xx,yy,sx=sxx,sy=syy),odr.Model(f),beta0=beta).run()
 # Independent joint least-squares optimization over coefficients and latent true abscissas.
 k=len(beta)
 def fun(v):return np.r_[(xx-v[k:])/sxx,(yy-f(v[:k],v[k:]))/syy]
 z=least_squares(fun,np.r_[o.beta,xx+o.delta],xtol=1e-13,ftol=1e-13,gtol=1e-13,max_nfev=10000)
 assert np.allclose(z.x[:k],o.beta,rtol=2e-5,atol=2e-6),(z.x[:k],o.beta)
 return {'beta':o.beta.tolist(),'u':np.sqrt(np.diag(o.cov_beta)).tolist(),'cov':o.cov_beta.tolist(),'chi2':float(o.sum_square),'dof':len(xx)-k,'reduced':float(o.res_var),'independent':z.x[:k].tolist(),'status':o.stopreason}
free=fit(X,Y,sx,sy,'line',[2.5,1.2])
quad=fit(Q,mean,uq,uy,'origin',[2.5])
quad_intercept=fit(Q,mean,uq,uy,'line',[2.5,-.03])
weights=1/(sy**2+4*sx**2);B=float(np.sum(weights*(Y-2*X))/weights.sum());uB=float(1/np.sqrt(weights.sum()))
logfix={'B':B,'uB':uB,'A':float(np.exp(B)),'uA':float(np.exp(B)*uB),'chi2':float(np.sum(weights*(Y-2*X-B)**2)),'dof':9}
logfix['reduced']=logfix['chi2']/9
for obj,A,uA in [(quad,quad['beta'][0],quad['u'][0]),(logfix,logfix['A'],logfix['uA']),(quad_intercept,quad_intercept['beta'][0],quad_intercept['u'][0])]:
 obj['v']=float(np.sqrt(g/(2*A)));obj['uv']=float(obj['v']*uA/(2*A))
free['C']=float(np.exp(free['beta'][1]));free['uC']=float(free['C']*free['u'][1])
mask=mean>3*uy
sensitivity=fit(X[mask],Y[mask],sx[mask],sy[mask],'line',free['beta'])
result={'date':'2026-10-13','g':g,'rows':[list(r[:4]) for r in rows],'x':x.tolist(),'y':mean.tolist(),'s':s.tolist(),'sem':sem.tolist(),'ux':ux.tolist(),'uy':uy.tolist(),'X':X.tolist(),'Y':Y.tolist(),'uX':sx.tolist(),'uY':sy.tolist(),'Q':Q.tolist(),'uQ':uq.tolist(),'power':free,'log_fixed':logfix,'quadratic':quad,'quadratic_intercept':quad_intercept,'log_sensitivity':{'included_x_cm':(100*x[mask]).tolist(),**sensitivity},'boundary':{'x_cm':55,'single_y_cm':81.7,'missing':2,'used_in_primary_fits':False}}
(Path(__file__).resolve().parent/'calculos_verificados.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({k:result[k] for k in ['power','log_fixed','quadratic','quadratic_intercept','log_sensitivity']},ensure_ascii=False,indent=2))
plt.rcParams.update({'font.family':'DejaVu Sans','font.size':11,'axes.spines.top':False,'axes.spines.right':False,'axes.grid':True,'grid.alpha':.22})
colors=['#204f72','#c56d20','#3b7863'];out=base/'graficos';out.mkdir(exist_ok=True)
def save(fig,name):
 fig.tight_layout();fig.savefig(out/(name+'.png'),dpi=240);fig.savefig(out/(name+'.pdf'));plt.close(fig)
def points(ax,xx,yy,xe,ye):ax.errorbar(xx,yy,xerr=xe,yerr=ye,fmt='o',markersize=4,capsize=2,color=colors[0],label='Médias de três impactos')
fig,ax=plt.subplots(figsize=(7.2,4.6));points(ax,x*100,mean*100,ux*100,uy*100);ax.plot(55,81.7,'D',mfc='none',color='#666666',label='55 cm: uma leitura, fora dos ajustes');ax.set(xlabel='Distância horizontal x (cm)',ylabel='Queda vertical y (cm)',title='Trajetória medida');ax.legend(fontsize=9);save(fig,'01_Trajetoria')
fig,ax=plt.subplots(figsize=(7.2,4.6));points(ax,Y,X,sy,sx);ax.set(xlabel='ln(y / 1 m)',ylabel='ln(x / 1 m)',title='');plt.close(fig)
fig,ax=plt.subplots(figsize=(7.2,4.6));points(ax,X,Y,sx,sy);grid=np.linspace(X.min()-.05,X.max()+.05,100);ax.plot(grid,free['beta'][0]*grid+free['beta'][1],color=colors[1],label='Expoente livre');ax.plot(grid,2*grid+B,'--',color=colors[2],label='Expoente fixo em 2');ax.set(xlabel='ln(x / 1 m)',ylabel='ln(y / 1 m)',title='Linearização logarítmica');ax.legend(fontsize=9);save(fig,'02_Linearizacao')
fig,ax=plt.subplots(figsize=(7.2,4.6));points(ax,x*100,mean*100,ux*100,uy*100);grid=np.linspace(.03,.51,200);ax.plot(grid*100,free['C']*grid**free['beta'][0]*100,color=colors[1],label='Lei de potência ajustada');ax.plot(grid*100,quad['beta'][0]*grid**2*100,'--',color=colors[2],label='Modelo quadrático');ax.set(xlabel='Distância horizontal x (cm)',ylabel='Queda vertical y (cm)',title='Comparação dos modelos');ax.legend(fontsize=9);save(fig,'03_Curvas_Ajustadas')
fig,ax=plt.subplots(figsize=(7.2,4.6));points(ax,Q,mean,uq,uy);grid=np.linspace(0,.27,100);ax.plot(grid,quad['beta'][0]*grid,color=colors[1],label='Ajuste pela origem');ax.set(xlabel='x² (m²)',ylabel='Queda vertical y (m)',title='Relação entre y e x²');ax.legend(fontsize=9);save(fig,'04_Ajuste_Quadratico')
fig,ax=plt.subplots(figsize=(7.2,4.6));points(ax,mean*100,x*100,uy*100,ux*100);ax.set(xlabel='Queda vertical y (cm)',ylabel='Distância horizontal x (cm)',title='Orientação solicitada no item 5 do roteiro');save(fig,'05_X_em_funcao_de_Y')

