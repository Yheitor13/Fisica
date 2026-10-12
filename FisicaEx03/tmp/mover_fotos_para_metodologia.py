from copy import deepcopy
from pathlib import Path
import re
from docx import Document
from docx.text.paragraph import Paragraph

base = Path(__file__).resolve().parents[1]
source = base / 'Relatorio_Final.docx'
target = base / 'tmp' / 'Relatorio_Metodologia_Candidato.docx'
doc = Document(source)

def index_of(paragraph):
    return next(i for i, item in enumerate(doc.paragraphs) if item._p is paragraph._p)

def unique(text):
    found = [p for p in doc.paragraphs if p.text == text]
    assert len(found) == 1, (text, len(found))
    return found[0]

method = next(p for p in doc.paragraphs if 'fotografias fornecidas pelo grupo documentam a montagem' in p.text)
assert len(method.runs) == 1 and 'no Apêndice A' in method.text
method.runs[0].text = method.runs[0].text.replace('no Apêndice A', 'na Seção 3.1')

result_heading = unique('4 RESULTADOS E DISCUSSÃO')
appendix_heading = unique('APÊNDICE A – REGISTRO FOTOGRÁFICO DA MONTAGEM')
start = index_of(appendix_heading)
blocks = doc.paragraphs[start:]
assert len(blocks) == 12 and blocks[-1].text == 'Fonte: registro fotográfico fornecido pelo grupo.'
for paragraph in blocks:
    result_heading._p.addprevious(paragraph._p)
appendix_heading.style = 'Heading 2'
appendix_heading.runs[0].text = '3.1 Registro fotográfico da montagem'
appendix_heading.paragraph_format.page_break_before = True
appendix_heading.paragraph_format.keep_with_next = True

figure_list_heading = unique('LISTA DE FIGURAS')
table_list_heading = unique('LISTA DE TABELAS')
fig_start = index_of(figure_list_heading) + 1
fig_end = index_of(table_list_heading)
figure_entries = doc.paragraphs[fig_start:fig_end]
assert len(figure_entries) == 8
photo_entries = figure_entries[5:]
first_graph_entry = figure_entries[0]
for paragraph in photo_entries:
    first_graph_entry._p.addprevious(paragraph._p)

mapping = {1:4, 2:5, 3:6, 4:7, 5:8, 6:1, 7:2, 8:3}
pattern = re.compile(r'\bFigura\s+([1-8])\b')
for paragraph in doc.paragraphs:
    for run in paragraph.runs:
        if pattern.search(run.text):
            run.text = pattern.sub(lambda m: 'Figura ' + str(mapping[int(m.group(1))]), run.text)

fig_start = index_of(figure_list_heading) + 1
fig_end = index_of(table_list_heading)
figure_entries = doc.paragraphs[fig_start:fig_end]
assert len(figure_entries) == 8
for paragraph, number, page in zip(figure_entries, range(1, 9), [8,9,9,11,12,13,14,15]):
    assert paragraph.text.startswith(f'Figura {number} – ')
    assert len(paragraph.runs) == 2
    paragraph.runs[1].text = '\t' + str(page)

table_entries = [p for p in doc.paragraphs if re.match(r'^Tabela [1-5] – .*\t\d+$', p.text)]
assert len(table_entries) == 5
for paragraph, page in zip(table_entries, [7,10,12,14,15]):
    assert len(paragraph.runs) == 2
    paragraph.runs[1].text = '\t' + str(page)

toc_appendix = unique('APÊNDICE A – REGISTRO FOTOGRÁFICO DA MONTAGEM\t17')
toc_appendix._p.getparent().remove(toc_appendix._p)
toc_method = unique('3 METODOLOGIA\t7')
toc_subsection_template = unique('4.1 Medidas e incertezas\t8')
element = deepcopy(toc_subsection_template._p)
toc_method._p.addnext(element)
toc_subsection = Paragraph(element, toc_method._parent)
assert len(toc_subsection.runs) == 2
toc_subsection.runs[0].text = '3.1 Registro fotográfico da montagem'
toc_subsection.runs[1].text = '\t8'

toc_pages = {
    '4 RESULTADOS E DISCUSSÃO':10,
    '4.1 Medidas e incertezas':10,
    '4.2 Lei de potência':12,
    '4.3 Dependência quadrática do tempo':14,
    '4.4 Interpretação dos resíduos':16,
    '5 CONCLUSÃO':17,
    'REFERÊNCIAS':18,
}
for label,page in toc_pages.items():
    matches=[p for p in doc.paragraphs if p.text.startswith(label+'\t')]
    assert len(matches)==1,(label,len(matches))
    paragraph=matches[0]
    assert len(paragraph.runs)==2
    paragraph.runs[1].text='\t'+str(page)

assert len(doc.inline_shapes)==8
assert not any('APÊNDICE A' in p.text for p in doc.paragraphs)
doc.save(target)
print(target)
