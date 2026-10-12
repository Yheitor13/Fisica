from copy import deepcopy
from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt
from docx.text.paragraph import Paragraph

base = Path(__file__).resolve().parents[1]
photos = base / 'refs' / 'Fotos'
assert all((photos / f'montagem_{idx:02}.jpeg').is_file() for idx in range(1, 6))

doc = Document(base / 'Relatorio_Final.docx')
assert not any('APÊNDICE A' in p.text for p in doc.paragraphs)
method = doc.paragraphs[69]
old = 'Esse é o arranjo prescrito pela apostila; o grupo forneceu as folhas de medições, sem fotografias da montagem, de modo que detalhes particulares do aparelho não são inferidos das imagens de anotações.'
new = 'As fotografias fornecidas pelo grupo documentam a montagem utilizada e são apresentadas no Apêndice A. Elas mostram a haste, o suporte da esfera, a base e a unidade de cronometragem, sem permitir concluir ajustes ou calibração que não foram registrados.'
assert method.text.count(old) == 1 and len(method.runs) == 1
method.runs[0].text = method.runs[0].text.replace(old, new)
ref = doc.paragraphs[126]
old = 'Planilha medidas_fisicaEX03.xlsx e duas fotografias das anotações. Conferido pelo grupo em 9 out. 2026.'
new = 'Planilha medidas_fisicaEX03.xlsx, duas fotografias das anotações e cinco fotografias da montagem. Medidas conferidas pelo grupo em 9 out. 2026; imagens da montagem fornecidas em 12 out. 2026.'
assert ref.text.count(old) == 1 and len(ref.runs) == 1
ref.runs[0].text = ref.runs[0].text.replace(old, new)

def add_list_line(after, label, page):
    element = deepcopy(after._p)
    after._p.addnext(element)
    paragraph = Paragraph(element, after._parent)
    assert len(paragraph.runs) == 2
    paragraph.runs[0].text = label
    paragraph.runs[1].text = '\t' + str(page)
    return paragraph

toc_reference = doc.paragraphs[47]
last = doc.paragraphs[26]
last = add_list_line(last, 'Figura 6 – Visão geral da montagem experimental', 17)
last = add_list_line(last, 'Figura 7 – Esfera, suporte e plataforma de parada', 18)
last = add_list_line(last, 'Figura 8 – Unidade de cronometragem', 18)
add_list_line(toc_reference, 'APÊNDICE A – REGISTRO FOTOGRÁFICO DA MONTAGEM', 17)

heading = doc.add_paragraph('APÊNDICE A – REGISTRO FOTOGRÁFICO DA MONTAGEM', style='Heading 1')
heading.paragraph_format.keep_with_next = True
intro = doc.add_paragraph('As imagens fornecidas pelo grupo registram o arranjo observado. Foram selecionadas três vistas complementares: o conjunto completo, a região de queda da esfera e a unidade de cronometragem.')
intro.paragraph_format.space_after = Pt(8)

def add_photo(number, caption, filename, width_cm):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(f'Figura {number} – {caption}')
    run.font.size = Pt(10)
    img = doc.add_paragraph()
    img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    img.paragraph_format.space_after = Pt(4)
    img.add_run().add_picture(str(photos / filename), width=Cm(width_cm))
    source = doc.add_paragraph()
    source.paragraph_format.space_after = Pt(3)
    s = source.add_run('Fonte: registro fotográfico fornecido pelo grupo.')
    s.font.size = Pt(9)

add_photo(6, 'Visão geral da montagem experimental', 'montagem_01.jpeg', 9.2)
doc.add_page_break()
add_photo(7, 'Esfera, suporte e plataforma de parada', 'montagem_04.jpeg', 5.7)
add_photo(8, 'Unidade de cronometragem', 'montagem_02.jpeg', 5.7)

target = base / 'tmp' / 'Relatorio_Fotos_Candidato.docx'
doc.save(target)
print(target)



